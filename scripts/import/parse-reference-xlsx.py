#!/usr/bin/env python3
"""
Convert the coach-provided reference spreadsheets into the normalized JSON
that the app ships (src/data/**). Re-run this if the source .xlsx files change.

Usage:
    pip install openpyxl
    python3 scripts/import/parse-reference-xlsx.py \
        --soccer path/to/Soccer_Reference_Information.xlsx \
        --swimming path/to/Swimming_Workout_Database.xlsx

Outputs:
    src/data/swimming/videos.json     video library keyed by name
    src/data/swimming/workouts.json   1,440 workouts (phases reference videos by name)
    src/data/swimming/index.json      lightweight list rows for browsing/filtering
    src/data/soccer/program.json      grade-level session plans + warm-up/cool-down
"""
import argparse
import json
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def parse_swimming(path: str) -> None:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    videos = {}
    for r in list(wb["Video Library"].iter_rows(values_only=True))[2:]:
        if not r or not r[0]:
            continue
        name, stroke, source, howto, skills, url = (list(r) + [None] * 6)[:6]
        videos[name] = {"stroke": stroke, "source": source, "howTo": howto, "skills": skills, "url": url}

    workouts = {}
    for r in list(wb["Workouts"].iter_rows(values_only=True))[2:]:
        if not r or not r[0]:
            continue
        wid, stroke, base, level, no, phase, setname, setdesc, dist, effort, howto, skills, source = (
            list(r) + [None] * 13
        )[:13]
        w = workouts.setdefault(
            wid, {"id": wid, "stroke": stroke, "base": base, "level": level, "no": no, "phases": []}
        )
        p = {"phase": phase, "video": setname, "set": setdesc, "distanceM": dist, "effort": effort}
        if setname not in videos:  # keep the row lossless if a video isn't in the library
            p["howTo"], p["skills"], p["source"] = howto, skills, source
        w["phases"].append(p)

    index = []
    for r in list(wb["Index"].iter_rows(values_only=True))[2:]:
        if not r or not r[0]:
            continue
        wid, stroke, base, level, no, target, total = (list(r) + [None] * 7)[:7]
        index.append(
            {"id": wid, "stroke": stroke, "base": base, "level": level, "no": no,
             "targetTime": target, "totalDistanceM": total}
        )
    wb.close()

    out = os.path.join(ROOT, "src/data/swimming")
    os.makedirs(out, exist_ok=True)
    for name, data in [("videos", videos), ("workouts", workouts), ("index", index)]:
        with open(os.path.join(out, f"{name}.json"), "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"swimming: {len(videos)} videos, {len(workouts)} workouts, {len(index)} index rows")


def _rows(wb, name):
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def parse_soccer(path: str) -> None:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def session(sheet, title):
        rows = _rows(wb, sheet)
        ncol = len(rows[1])
        cols = [[] for _ in range(ncol)]
        for r in rows[3:]:
            for ci, val in enumerate(r):
                if val not in (None, ""):
                    cols[ci].append(str(val).strip())
        phases = []
        for ci in range(1, ncol):
            head = rows[1][ci]
            if not head:
                continue
            time = rows[2][ci] if ci < len(rows[2]) else None
            phases.append({"name": str(head).strip(), "time": str(time).strip() if time else None, "items": cols[ci]})
        return {"title": title, "phases": phases}

    def listing(sheet):
        rows = _rows(wb, sheet)
        start = 1 if rows and rows[0] and str(rows[0][0]).lower() == "name" else 0
        out = []
        for r in rows[start:]:
            vals = [str(c).strip() for c in r if c not in (None, "")]
            if vals:
                out.append({"name": vals[0], "variations": vals[1:]})
        return out

    data = {
        "sessions": [
            session("Pre-First - Grade 2", "유치원~2학년"),
            session("Grade 3 - 5", "3~5학년"),
        ],
        "warmup": listing("Stretching Routine"),
        "cooldown": listing("Cool-down stretch"),
    }
    wb.close()

    out = os.path.join(ROOT, "src/data/soccer")
    os.makedirs(out, exist_ok=True)
    with open(os.path.join(out, "program.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"soccer: {len(data['sessions'])} sessions, {len(data['warmup'])} warm-up, {len(data['cooldown'])} cool-down")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--soccer")
    ap.add_argument("--swimming")
    args = ap.parse_args()
    if args.swimming:
        parse_swimming(args.swimming)
    if args.soccer:
        parse_soccer(args.soccer)
